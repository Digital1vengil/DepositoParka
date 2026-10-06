"""Sección Control de paquetes: escanear la etiqueta de envío y cada prenda; la app verifica que coincida
con la venta y no deja cerrar el paquete si algo está mal."""
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..actividad import registrar
from ..auth import puede, requiere, requiere_alguno
from ..db import get_db, hoy
from ..models import Articulo, Control, ControlEvento, ControlItem, ControlPaquete, Usuario
from ..services import control as svc
from ..services import devoluciones as dsvc
from ..services import drive
from ..templating import render

router = APIRouter(prefix="/control")
acceso = requiere_alguno("control", "ecommerce")  # la mesa de control de Recolección Ecommerce usa esta API
TIPOS = {"flex": "Flex", "colecta": "Colecta", "tiendanube": "Tienda Nube"}
MAX_BYTES = 15 * 1024 * 1024


def _tipos(tipo: str) -> list[str]:
    """"flex,colecta" → ["flex", "colecta"] (se pueden elegir varios tipos juntos)."""
    return [t for t in (tipo or "").split(",") if t in TIPOS]


def _tipo_parse(tipo: str) -> str:
    ts = _tipos(tipo)
    return ts[0] if len(ts) == 1 else ""


def _tipo_guardar(tipo: str) -> str:
    ts = _tipos(tipo)
    return ts[0] if len(ts) == 1 else ("mixto" if ts else (tipo or "")[:20])


def _nombre_tipos(tipo: str) -> str:
    return " + ".join(TIPOS[t] for t in _tipos(tipo)) or "Ventas"
ESTADOS = {"pendiente": "Pendiente", "en_curso": "En curso", "completo": "Controlado", "faltante": "Con faltante"}


@router.get("")
def pagina(request: Request, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    re_id = request.query_params.get("re", "")
    return render(request, "control.html", u=u, drive_ok=drive.configurado(),
                  puede_vincular=puede(db, u, "supervision"), re_id=int(re_id) if re_id.isdigit() else 0)


def _control(db: Session, cid: int) -> Control:
    c = db.get(Control, cid)
    if not c:
        raise HTTPException(404, "No existe ese control")
    return c


def _paquete(db: Session, pid: int) -> ControlPaquete:
    p = db.get(ControlPaquete, pid)
    if not p:
        raise HTTPException(404, "No existe ese paquete")
    return p


def _item(db: Session, iid: int) -> ControlItem:
    i = db.get(ControlItem, iid)
    if not i:
        raise HTTPException(404, "No existe esa línea")
    return i


def _abierto(c: Control) -> None:
    if c.estado != "abierto":
        raise HTTPException(400, "Este control ya está cerrado (reabrilo para seguir)")


async def _leer_subidos(archivos: list[UploadFile]) -> list[tuple[str, bytes]]:
    datos = []
    for f in archivos:
        b = await f.read()
        if len(b) > MAX_BYTES:
            raise HTTPException(400, f"{f.filename}: archivo demasiado grande")
        datos.append((f.filename or "archivo", b))
    return datos


def _crear(db: Session, u: Usuario, paquetes: list[dict], origen: str, tipo: str, nombre: str,
           errores: list[str], archivos: str = "") -> dict:
    if not paquetes:
        raise HTTPException(400, " · ".join(errores) or "No hay paquetes para controlar")
    nombre = nombre.strip() or f"{_nombre_tipos(tipo)} {hoy().strftime('%d/%m')}"
    tipo = _tipo_guardar(tipo)
    c = svc.crear(db, u, paquetes, origen, tipo, nombre, archivos)
    registrar(db, u, "control_creado", f"CP-{c.id} · {c.nombre} · {len(c.paquetes)} paquetes", "control")
    return {"ok": True, "control": svc.detalle(c), "errores": errores}


@router.post("/api/crear/archivo")
async def crear_archivo(tipo: str = Form("flex"), nombre: str = Form(""), archivos: list[UploadFile] = File(...),
                        u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    datos = await _leer_subidos(archivos)
    paquetes, errores = svc.leer_archivos(datos, _tipo_parse(tipo))
    return _crear(db, u, paquetes, "archivo", tipo, nombre, errores, ", ".join(n for n, _ in datos))


@router.get("/api/drive/archivos")
def drive_archivos(tipo: str = "flex", u: Usuario = Depends(acceso)):
    try:
        d = drive.listar(tipo)
    except drive.DriveError as e:
        raise HTTPException(502, str(e))
    return {"grupos": (d.get("groups") or [])[:10], "hoy": hoy().isoformat()}


class DriveIn(BaseModel):
    ids: list[str]
    tipo: str = "flex"
    nombre: str = ""


@router.post("/api/crear/drive")
def crear_drive(body: DriveIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    if not body.ids:
        raise HTTPException(400, "Elegí al menos un archivo")
    try:
        datos = [drive.obtener_archivo(i) for i in body.ids[:20]]
    except drive.DriveError as e:
        raise HTTPException(502, str(e))
    paquetes, errores = svc.leer_archivos(datos, _tipo_parse(body.tipo))
    return _crear(db, u, paquetes, "drive", body.tipo, body.nombre, errores, ", ".join(n for n, _ in datos))


class PaqueteIn(BaseModel):
    tracking: str = ""
    venta_id: str = ""
    comprador: str = ""
    producto: str = ""
    sku: str = ""
    cantidad: int = 1


class PaquetesIn(BaseModel):
    paquetes: list[PaqueteIn]
    tipo: str = ""
    nombre: str = ""


@router.post("/api/crear/paquetes")
def crear_paquetes(body: PaquetesIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Desde la lista cargada en Despacho (el scanner) en este dispositivo."""
    return _crear(db, u, [p.model_dump() for p in body.paquetes], "despacho", body.tipo, body.nombre, [])


@router.post("/api/{cid}/agregar")
async def agregar(cid: int, archivos: list[UploadFile] = File(...), u: Usuario = Depends(acceso),
                  db: Session = Depends(get_db)):
    """Sumar archivos a un control (ej. la planilla de ventas después del TXT de etiquetas)."""
    c = _control(db, cid)
    _abierto(c)
    datos = await _leer_subidos(archivos)
    base = [dict(tracking=p.tracking, tracking_alt=p.tracking_alt, venta_id=p.venta_id, comprador=p.comprador,
                 unidades=p.unidades_etiqueta,
                 items=[dict(sku=i.sku, producto=i.producto, variante=i.variante, cantidad=i.cantidad) for i in p.items])
            for p in c.paquetes]
    paquetes, errores = svc.leer_archivos(datos, c.tipo, base=base)
    nuevos, act = svc.agregar_paquetes(db, c, paquetes)
    c.archivos = (((c.archivos + ", ") if c.archivos else "") + ", ".join(n for n, _ in datos))[:2000]
    db.commit()
    return {"ok": True, "nuevos": nuevos, "actualizados": act, "errores": errores, "control": svc.detalle(c)}


@router.get("/api/lista")
def lista(u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    cs = db.scalars(select(Control).order_by(Control.id.desc()).limit(20)).all()
    return {"controles": [svc.resumen(c) for c in cs]}


@router.get("/api/buscar")
def buscar(q: str = "", todos: bool = False, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"paquetes": [svc.paquete_json(p, con_items=False) | {"control": p.control.nombre}
                         for p in svc.buscar(db, q, solo_hoy=not todos)]}


@router.get("/api/catalogo")
def catalogo(q: str = "", u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"modelos": dsvc.buscar_modelos(db, q, filtro=Articulo.id.is_not(None), limite=20)}


@router.get("/api/estado")
def estado(codigo: str, u: Usuario = Depends(requiere("despacho")), db: Session = Depends(get_db)):
    """Para Despacho: ¿este envío se controló hoy?"""
    from ..parsers import codigo_comparable, extraer_id
    cid = codigo_comparable(extraer_id(codigo))
    ps = db.scalars(select(ControlPaquete).join(Control).where(Control.fecha == hoy())).all()
    p = next((x for x in ps if cid and cid in (codigo_comparable(x.tracking), codigo_comparable(x.tracking_alt))), None)
    if not p:
        return {"en_control": False}
    return {"en_control": True, "estado": p.estado, "estado_txt": ESTADOS.get(p.estado, p.estado),
            "controlado": p.estado == "completo", "comprador": p.comprador, "control": p.control.nombre}


@router.get("/api/{cid}")
def ver(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"control": svc.detalle(_control(db, cid))}


@router.get("/api/paquete/{pid}")
def ver_paquete(pid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    p = _paquete(db, pid)
    return {"paquete": svc.paquete_json(p), "resumen": svc.resumen(p.control)}


class EscanearIn(BaseModel):
    codigo: str
    paquete_id: int | None = None


@router.post("/api/{cid}/escanear")
def escanear(cid: int, body: EscanearIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Un solo campo para todo: etiqueta de envío → abre el paquete; prenda → la compara con el paquete abierto."""
    c = _control(db, cid)
    _abierto(c)
    cod = body.codigo.strip()
    if not cod:
        raise HTTPException(400, "Código vacío")
    p = svc.buscar_paquete(c, cod)
    if p:
        anterior = db.get(ControlPaquete, body.paquete_id) if body.paquete_id else None
        aviso = ""
        if anterior and anterior.id != p.id and anterior.estado == "en_curso":
            aviso = f"El paquete {anterior.tracking} quedó a medias ({svc.unidades(anterior)[0]}/{svc.unidades(anterior)[1]})"
        return {"tipo": "paquete", "paquete": svc.paquete_json(p), "aviso": aviso, "resumen": svc.resumen(c)}
    actual = db.get(ControlPaquete, body.paquete_id) if body.paquete_id else None
    if actual and actual.control_id == c.id:
        if actual.estado == "completo":
            return {"tipo": "prenda", "ok": False, "motivo": "cerrado", "paquete": svc.paquete_json(actual),
                    "mensaje": "Este paquete ya está controlado. Escaneá la etiqueta del próximo envío."}
        r = svc.escanear_prenda(db, actual, u, cod)
        if not r["ok"]:
            registrar(db, u, "control_error", f"{actual.tracking}: {r.get('escaneado') or cod}", "control")
        elif r.get("completo"):
            registrar(db, u, "control_completo", f"{actual.tracking} · {actual.comprador}", "control")
        return {"tipo": "prenda", **r, "paquete": svc.paquete_json(actual), "resumen": svc.resumen(c)}
    # sin paquete abierto: ¿es una prenda? → decir en qué paquete va
    art = svc.articulo_de_codigo(db, cod)
    if art:
        ps = svc.paquetes_que_llevan(c, art)
        return {"tipo": "guia", "escaneado": svc.nombre_art(art),
                "paquetes": [svc.paquete_json(x, con_items=False) for x in ps],
                "mensaje": (f"{svc.nombre_art(art)} va en {len(ps)} paquete(s): escaneá la etiqueta del envío"
                            if ps else f"{svc.nombre_art(art)} no está en ningún paquete pendiente de este control")}
    otro = db.scalars(select(ControlPaquete).join(Control).where(Control.fecha == hoy(), Control.id != c.id)).all()
    from ..parsers import codigo_comparable, extraer_id
    cid_ = codigo_comparable(extraer_id(cod))
    o = next((x for x in otro if cid_ and cid_ == codigo_comparable(x.tracking)), None)
    if o:
        raise HTTPException(404, f"Ese envío está en otro control de hoy: {o.control.nombre} (CP-{o.control_id})")
    raise HTTPException(404, "No es un envío de este control ni un artículo del catálogo")


class FaltanteIn(BaseModel):
    nota: str = ""


@router.post("/api/item/{iid}/faltante")
def faltante(iid: int, body: FaltanteIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    i = _item(db, iid)
    _abierto(i.paquete.control)
    svc.marcar_faltante(db, i, u, body.nota.strip())
    registrar(db, u, "control_faltante", f"{i.paquete.tracking}: {i.articulo or i.producto[:40]}", "control")
    return {"ok": True, "paquete": svc.paquete_json(i.paquete), "resumen": svc.resumen(i.paquete.control)}


@router.post("/api/item/{iid}/quitar-faltante")
def quitar_faltante(iid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    i = _item(db, iid)
    _abierto(i.paquete.control)
    svc.quitar_faltante(db, i, u)
    return {"ok": True, "paquete": svc.paquete_json(i.paquete), "resumen": svc.resumen(i.paquete.control)}


@router.get("/api/item/{iid}/donde")
def donde(iid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    i = _item(db, iid)
    return {"opciones": svc.donde_hay(db, i)}


@router.post("/api/paquete/{pid}/reiniciar")
def reiniciar(pid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    p = _paquete(db, pid)
    _abierto(p.control)
    svc.reiniciar(db, p, u)
    registrar(db, u, "control_reinicio", p.tracking, "control")
    return {"ok": True, "paquete": svc.paquete_json(p), "resumen": svc.resumen(p.control)}


class VincularIn(BaseModel):
    articulo_id: int


@router.post("/api/item/{iid}/vincular")
def vincular(iid: int, body: VincularIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Asigna el artículo del catálogo al SKU de la venta. Solo admin o quien tenga Supervisión,
    para que nadie 'arregle' un error de armado vinculando la prenda equivocada."""
    if not puede(db, u, "supervision"):
        raise HTTPException(403, "Vincular artículos lo hace un supervisor (admin o con acceso a Supervisión)")
    i, a = _item(db, iid), db.get(Articulo, body.articulo_id)
    if not a:
        raise HTTPException(404, "No existe ese artículo")
    n = svc.vincular(db, i, a, u)
    registrar(db, u, "sku_vinculado", f"{i.sku or i.producto[:40]} → {a.nombre}", "control")
    return {"ok": True, "lineas": n, "paquete": svc.paquete_json(i.paquete), "resumen": svc.resumen(i.paquete.control)}


@router.get("/api/{cid}/eventos")
def eventos(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    evs = db.scalars(select(ControlEvento).where(ControlEvento.control_id == cid)
                     .order_by(ControlEvento.id.desc()).limit(200)).all()
    trk = {p.id: p.tracking for p in _control(db, cid).paquetes}
    return {"eventos": [{"hora": e.ts.strftime("%H:%M"), "tipo": e.tipo, "tracking": trk.get(e.paquete_id, ""),
                         "usuario": e.usuario.nombre if e.usuario else "", "detalle": e.detalle, "codigo": e.codigo}
                        for e in evs]}


@router.post("/api/{cid}/cerrar")
def cerrar(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    from ..db import ahora
    c = _control(db, cid)
    c.estado, c.cerrado_at = "cerrado", ahora()
    db.commit()
    return {"ok": True, "control": svc.detalle(c)}


@router.post("/api/{cid}/reabrir")
def reabrir(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    c = _control(db, cid)
    c.estado, c.cerrado_at = "abierto", None
    db.commit()
    return {"ok": True, "control": svc.detalle(c)}


@router.delete("/api/{cid}")
def borrar(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    c = _control(db, cid)
    if not u.es_admin and c.usuario_id != u.id:
        raise HTTPException(403, "Solo quien lo cargó (o un admin) puede borrarlo")
    db.execute(delete(ControlEvento).where(ControlEvento.control_id == cid))
    db.delete(c)
    db.commit()
    return {"ok": True}


@router.get("/{cid}/control.xlsx")
def excel(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    c = _control(db, cid)
    s = svc.resumen(c)
    wb = Workbook()
    ws = wb.active
    ws.title = "Paquetes"
    ws.append([f"Control CP-{c.id} · {c.nombre}", c.creado_at.strftime("%d/%m/%Y %H:%M"),
               f"Controlados {s['completo']}/{s['paquetes']}", f"Errores {s['errores']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Envío", "Venta", "Comprador", "Estado", "Artículo", "Color", "Talle", "Cant.", "Escaneadas",
               "Faltante", "Ubicación", "SKU venta", "Publicación", "Errores", "Controló", "Hora", "Nota"])
    for cel in ws[3]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="3B1F6B")
    rojo = Font(bold=True, color="C8283E")
    for p in c.paquetes:
        for i in p.items:
            ws.append([p.tracking, p.venta_id, p.comprador, ESTADOS.get(p.estado, p.estado), i.articulo or "SIN VINCULAR",
                       i.color, i.talle, i.cantidad, i.escaneado, "SÍ" if i.faltante else "", i.ubicacion, i.sku,
                       i.producto, p.errores, p.controlado_por.nombre if p.controlado_por else "",
                       p.controlado_at.strftime("%H:%M") if p.controlado_at else "", p.nota])
            if p.estado != "completo":
                ws.cell(ws.max_row, 4).font = rojo
    for col, ancho in zip("ABCDEFGHIJKLMNOPQ", (14, 18, 24, 13, 22, 16, 7, 6, 10, 9, 14, 22, 40, 8, 12, 7, 30)):
        ws.column_dimensions[col].width = ancho
    ws2 = wb.create_sheet("Errores y faltantes")
    ws2.append(["Hora", "Envío", "Tipo", "Usuario", "Detalle", "Código"])
    trk = {p.id: p.tracking for p in c.paquetes}
    for e in db.scalars(select(ControlEvento).where(ControlEvento.control_id == cid).order_by(ControlEvento.id)).all():
        ws2.append([e.ts.strftime("%H:%M"), trk.get(e.paquete_id, ""), e.tipo, e.usuario.nombre if e.usuario else "",
                    e.detalle, e.codigo])
    for col, ancho in zip("ABCDEF", (7, 14, 12, 12, 70, 16)):
        ws2.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="control_CP-{c.id}.xlsx"'})

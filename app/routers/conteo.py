"""Sección Conteo: se fija un artículo, se cargan cantidades por color y talle, y sale un Excel (y etiquetas opcionales)."""
import io

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..actividad import registrar
from ..auth import requiere
from ..db import get_db, hoy
from ..models import Articulo, Conteo, ConteoItem, Usuario
from ..services import devoluciones as svc
from ..services.etiquetas import pdf_etiquetas
from ..services.etiquetas_html import html_etiquetas
from ..templating import render

router = APIRouter(prefix="/conteo")
acceso = requiere("conteo")
MARCAS_CONTEO = ("Parka", "Puffers", "No Brand")


@router.get("")
def pagina(request: Request, u: Usuario = Depends(acceso)):
    return render(request, "conteo.html", u=u)


@router.get("/api/buscar")
def buscar(q: str = "", u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"modelos": svc.buscar_modelos(db, q, filtro=svc.es_conteo())}


class ItemIn(BaseModel):
    articulo_id: int
    cantidad: int


class ConfirmarIn(BaseModel):
    items: list[ItemIn]
    nota: str = ""


def conteo_json(c: Conteo) -> dict:
    return {"id": c.id, "hora": c.creado_at.strftime("%d/%m %H:%M"), "total": c.total, "nota": c.nota,
            "usuario": c.usuario.nombre if c.usuario else "",
            "items": [{"articulo": i.articulo, "color": i.color, "talle": i.talle, "marca": i.marca,
                       "cantidad": i.cantidad, "sku": i.sku} for i in c.items]}


@router.post("/api/confirmar")
def confirmar(body: ConfirmarIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    cantidades: dict[int, int] = {}
    for it in body.items:
        if it.cantidad > 0:
            cantidades[it.articulo_id] = cantidades.get(it.articulo_id, 0) + it.cantidad
    if not cantidades:
        raise HTTPException(400, "Cargá al menos una cantidad")
    c = Conteo(fecha=hoy(), usuario_id=u.id, nota=body.nota.strip()[:200])
    db.add(c)
    for aid, cant in cantidades.items():
        a = db.get(Articulo, aid)
        if not a or not a.codigo_barras:
            raise HTTPException(400, "Hay un artículo que no existe o no tiene código de barras")
        m = svc.marca(a)
        if m not in MARCAS_CONTEO:
            raise HTTPException(400, f"{a.nombre}: el conteo es solo de artículos Parka, Puffers o No Brand")
        color, talle = svc.color_talle(a)
        c.items.append(ConteoItem(articulo_id=a.id, codigo_barras=a.codigo_barras, sku=a.sku, nombre=a.nombre,
                                  articulo=a.articulo, color=color, talle=talle, marca=m,
                                  cantidad=min(cant, 99999)))
    c.total = sum(i.cantidad for i in c.items)
    db.commit()
    registrar(db, u, "conteo", f"CT-{c.id} · {c.total} unidades", "conteo")
    return {"ok": True, "conteo": conteo_json(c)}


@router.get("/api/historial")
def historial(u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    cs = db.scalars(select(Conteo).order_by(Conteo.id.desc()).limit(25)).all()
    return {"conteos": [conteo_json(c) for c in cs]}


def _conteo(db: Session, cid: int) -> Conteo:
    c = db.get(Conteo, cid)
    if not c:
        raise HTTPException(404)
    return c


@router.get("/{cid}/conteo.xlsx")
def excel(cid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    c = _conteo(db, cid)
    wb = Workbook()
    ws = wb.active
    ws.title = "Conteo"
    ws.append([f"Conteo CT-{c.id}", c.creado_at.strftime("%d/%m/%Y %H:%M"),
               f"Responsable: {c.usuario.nombre if c.usuario else ''}", c.nota])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    enc = ["Artículo", "Color", "Talle", "Marca", "SKU", "Código de barras", "Cantidad"]
    ws.append(enc)
    for cel in ws[3]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="3B1F6B")
    for i in sorted(c.items, key=lambda x: (x.articulo, x.color, x.id)):
        ws.append([i.articulo, i.color, i.talle, i.marca, i.sku, i.codigo_barras, i.cantidad])
    ws.append(["", "", "", "", "", "TOTAL", c.total])
    ws.cell(ws.max_row, 6).font = ws.cell(ws.max_row, 7).font = Font(bold=True)
    for col, ancho in zip("ABCDEFG", (22, 20, 8, 11, 34, 17, 10)):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    nombre = f"conteo_CT-{c.id}_{c.creado_at.strftime('%Y-%m-%d')}.xlsx"
    return Response(buf.getvalue(), headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/{cid}/etiquetas.html")
def etiquetas_html(cid: int, desde: int = 1, guias: bool = False, u: Usuario = Depends(acceso),
                   db: Session = Depends(get_db)):
    from fastapi.responses import HTMLResponse
    c = _conteo(db, cid)
    lista = [{"nombre": i.nombre, "articulo": i.articulo, "color": i.color, "talle": i.talle,
              "sku": i.sku, "codigo_barras": i.codigo_barras} for i in c.items for _ in range(i.cantidad)]
    return HTMLResponse(html_etiquetas(lista, desde=desde, guias=guias, titulo=f"Etiquetas CT-{c.id}"))


@router.get("/{cid}/etiquetas.pdf")
def etiquetas(cid: int, desde: int = 1, guias: bool = False, u: Usuario = Depends(acceso),
              db: Session = Depends(get_db)):
    c = _conteo(db, cid)
    lista = [{"nombre": i.nombre, "articulo": i.articulo, "color": i.color, "talle": i.talle,
              "sku": i.sku, "codigo_barras": i.codigo_barras} for i in c.items for _ in range(i.cantidad)]
    return Response(pdf_etiquetas(lista, desde=desde, guias=guias), media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="etiquetas_CT-{c.id}.pdf"'})

"""Sección Recolección: lista de artículos a juntar para las ventas del día, en el orden de la ruta del depósito."""
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..actividad import registrar
from ..auth import requiere
from ..db import ahora, get_db, hoy
from ..models import Articulo, Recoleccion, RecoleccionItem, Ubicacion, Usuario
from ..parsers import leer_archivo
from ..services import devoluciones as dsvc
from ..services import drive, mapeo
from ..services import recoleccion as svc
from ..templating import render

router = APIRouter(prefix="/recoleccion")
acceso = requiere("recoleccion")
TIPOS = {"flex": "Flex", "colecta": "Colecta", "tiendanube": "Tienda Nube"}
MAX_BYTES = 15 * 1024 * 1024


def _tipos(tipo: str) -> list[str]:
    """"flex,colecta" → ["flex", "colecta"] (se pueden elegir varios tipos juntos)."""
    return [t for t in (tipo or "").split(",") if t in TIPOS]


def _tipo_parse(tipo: str) -> str:
    ts = _tipos(tipo)
    return ts[0] if len(ts) == 1 else ""


def _tipo_guardar(tipo: str) -> str:
    ts = _tipos(tipo)
    return ts[0] if len(ts) == 1 else ("mixto" if ts else (tipo or "")[:20])


def _nombre_tipos(tipo: str) -> str:
    return " + ".join(TIPOS[t] for t in _tipos(tipo)) or "Ventas"


@router.get("")
def pagina(request: Request, u: Usuario = Depends(acceso)):
    return render(request, "recoleccion.html", u=u, drive_ok=drive.configurado())


def _paquetes_de_archivos(archivos: list[tuple[str, bytes]], tipo: str) -> tuple[list[dict], list[str]]:
    paquetes, errores = [], []
    for nombre, data in archivos:
        try:
            r = leer_archivo(nombre, data, tipo, None)
        except ValueError as e:
            errores.append(f"{nombre}: {e}")
            continue
        if not r["paquetes"]:
            errores.append(f"{nombre}: no se encontraron paquetes")
        paquetes.extend(r["paquetes"])
    vistos, unicos = set(), []
    for p in paquetes:
        k = p.get("tracking") or id(p)
        if k not in vistos:
            vistos.add(k)
            unicos.append(p)
    return unicos, errores


def _crear(db: Session, u: Usuario, paquetes: list[dict], origen: str, tipo: str, nombre: str,
           errores: list[str] | None = None) -> dict:
    if not paquetes:
        raise HTTPException(400, " · ".join(errores or []) or "No hay paquetes para recolectar")
    nombre = nombre.strip() or f"{_nombre_tipos(tipo)} {hoy().strftime('%d/%m')}"
    tipo = _tipo_guardar(tipo)
    r = svc.crear(db, u, paquetes, origen, tipo, nombre)
    registrar(db, u, "recoleccion_creada", f"RC-{r.id} · {r.nombre} · {len(r.items)} líneas", "recoleccion")
    return {"ok": True, "recoleccion": svc.detalle(r), "errores": errores or []}


@router.get("/api/drive/archivos")
def drive_archivos(tipo: str = "flex", u: Usuario = Depends(acceso)):
    try:
        d = drive.listar(tipo)
    except drive.DriveError as e:
        raise HTTPException(502, str(e))
    return {"grupos": (d.get("groups") or [])[:10], "hoy": hoy().isoformat()}


class DriveIn(BaseModel):
    ids: list[str]
    tipo: str = "flex"
    nombre: str = ""


@router.post("/api/crear/drive")
def crear_drive(body: DriveIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    if not body.ids:
        raise HTTPException(400, "Elegí al menos un archivo")
    try:
        archivos = [drive.obtener_archivo(i) for i in body.ids[:20]]
    except drive.DriveError as e:
        raise HTTPException(502, str(e))
    paquetes, errores = _paquetes_de_archivos(archivos, _tipo_parse(body.tipo))
    return _crear(db, u, paquetes, "drive", body.tipo, body.nombre, errores)


class PaqueteIn(BaseModel):
    tracking: str = ""
    venta_id: str = ""
    producto: str = ""
    sku: str = ""
    cantidad: int = 1


class PaquetesIn(BaseModel):
    paquetes: list[PaqueteIn]
    tipo: str = ""
    nombre: str = ""


@router.post("/api/crear/paquetes")
def crear_paquetes(body: PaquetesIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Desde la lista que está cargada en Despacho (el scanner) en este dispositivo."""
    return _crear(db, u, [p.model_dump() for p in body.paquetes], "despacho", body.tipo, body.nombre)


@router.post("/api/crear/archivo")
async def crear_archivo(tipo: str = Form("flex"), nombre: str = Form(""), archivos: list[UploadFile] = File(...),
                        u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    datos = []
    for f in archivos:
        b = await f.read()
        if len(b) > MAX_BYTES:
            raise HTTPException(400, f"{f.filename}: archivo demasiado grande")
        datos.append((f.filename or "archivo", b))
    paquetes, errores = _paquetes_de_archivos(datos, _tipo_parse(tipo))
    return _crear(db, u, paquetes, "archivo", tipo, nombre, errores)


@router.get("/api/buscar")
def buscar(q: str = "", u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"modelos": dsvc.buscar_modelos(db, q, filtro=Articulo.id.is_not(None), limite=20)}


@router.get("/api/lista")
def lista(u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    rs = db.scalars(select(Recoleccion).order_by(Recoleccion.id.desc()).limit(20)).all()
    return {"recolecciones": [svc.resumen(r) for r in rs]}


def _rec(db: Session, rid: int) -> Recoleccion:
    r = db.get(Recoleccion, rid)
    if not r:
        raise HTTPException(404, "No existe esa recolección")
    return r


def _abierta(r: Recoleccion) -> None:
    if r.estado != "abierta":
        raise HTTPException(400, "Esta recolección ya está cerrada")


@router.get("/api/{rid}")
def ver(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    return {"recoleccion": svc.detalle(_rec(db, rid))}


class CodigoIn(BaseModel):
    codigo: str


@router.post("/api/{rid}/escanear")
def escanear(rid: int, body: CodigoIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    r = _rec(db, rid)
    _abierta(r)
    cod = body.codigo.strip()
    candidatos = [i for i in svc.orden_items(r) if i.codigo_barras and i.codigo_barras == cod]
    if not candidatos:
        a = mapeo.articulo_por_sku(db, cod)
        if a:
            candidatos = [i for i in svc.orden_items(r) if i.articulo_id == a.id]
    if not candidatos:
        ub = db.scalar(select(Ubicacion).where(Ubicacion.codigo == mapeo.normalizar_codigo(cod)))
        if ub:
            aca = [svc.item_json(i) for i in svc.orden_items(r)
                   if ub.codigo in [c.strip() for c in i.ubicacion.split(",")] and i.recogido < i.cantidad]
            return {"ok": True, "tipo": "ubicacion", "ubicacion": ub.codigo, "items": aca}
        raise HTTPException(404, "Ese artículo no está en esta recolección")
    pendiente = next((i for i in candidatos if i.recogido < i.cantidad), None)
    if not pendiente:
        raise HTTPException(409, f"{candidatos[0].articulo} {candidatos[0].color} {candidatos[0].talle}: ya está completo")
    svc.marcar(pendiente, 1, u)
    db.commit()
    return {"ok": True, "tipo": "item", "item": svc.item_json(pendiente), "resumen": svc.resumen(r)}


class MarcarIn(BaseModel):
    delta: int = 1
    todo: bool = False


@router.post("/api/item/{iid}/marcar")
def marcar(iid: int, body: MarcarIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    i = db.get(RecoleccionItem, iid)
    if not i:
        raise HTTPException(404)
    _abierta(i.recoleccion)
    svc.marcar(i, (i.cantidad - i.recogido) if body.todo else body.delta, u)
    db.commit()
    return {"ok": True, "item": svc.item_json(i), "resumen": svc.resumen(i.recoleccion)}


class VincularIn(BaseModel):
    articulo_id: int
    recordar: bool = True


@router.post("/api/item/{iid}/vincular")
def vincular(iid: int, body: VincularIn, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Asigna el artículo del catálogo a una línea; si tiene SKU, queda recordado para las próximas."""
    i = db.get(RecoleccionItem, iid)
    a = db.get(Articulo, body.articulo_id)
    if not i or not a:
        raise HTTPException(404)
    r = i.recoleccion
    _abierta(r)
    afectados = [x for x in r.items if i.sku and x.sku.upper() == i.sku.upper()] or [i]
    for x in afectados:
        svc.aplicar_articulo(db, x, a)
    if body.recordar and i.sku:
        mapeo.vincular_sku(db, i.sku, a.id, u.id)
    db.commit()
    registrar(db, u, "sku_vinculado", f"{i.sku or i.producto[:40]} → {a.nombre}", "recoleccion")
    return {"ok": True, "recoleccion": svc.detalle(r)}


@router.post("/api/{rid}/actualizar")
def actualizar(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    """Vuelve a buscar las ubicaciones (por si se cargaron en Mapeo después de armar la lista)."""
    r = _rec(db, rid)
    for it in r.items:
        if not it.articulo_id:
            a = mapeo.articulo_por_sku(db, it.sku)
            if a:
                svc.aplicar_articulo(db, it, a)
    svc.recalcular_ubicaciones(db, r)
    return {"ok": True, "recoleccion": svc.detalle(r)}


@router.post("/api/{rid}/cerrar")
def cerrar(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    r = _rec(db, rid)
    _abierta(r)
    r.estado, r.cerrada_at, r.cerrada_por_id = "cerrada", ahora(), u.id
    db.commit()
    s = svc.resumen(r)
    registrar(db, u, "recoleccion_cerrada", f"RC-{r.id} · {s['recogidas']}/{s['unidades']} u. · faltan {s['faltan']}",
              "recoleccion")
    return {"ok": True, "recoleccion": svc.detalle(r)}


@router.post("/api/{rid}/reabrir")
def reabrir(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    r = _rec(db, rid)
    r.estado, r.cerrada_at, r.cerrada_por_id = "abierta", None, None
    db.commit()
    return {"ok": True, "recoleccion": svc.detalle(r)}


@router.delete("/api/{rid}")
def borrar(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    r = _rec(db, rid)
    if not u.es_admin and r.usuario_id != u.id:
        raise HTTPException(403, "Solo quien la armó (o un admin) puede borrarla")
    db.delete(r)
    db.commit()
    registrar(db, u, "recoleccion_borrada", f"RC-{rid}", "recoleccion")
    return {"ok": True}


@router.get("/{rid}/recoleccion.xlsx")
def excel(rid: int, u: Usuario = Depends(acceso), db: Session = Depends(get_db)):
    r = _rec(db, rid)
    s = svc.resumen(r)
    wb = Workbook()
    ws = wb.active
    ws.title = "Recolección"
    ws.append([f"Recolección RC-{r.id} · {r.nombre}", r.creado_at.strftime("%d/%m/%Y %H:%M"),
               f"Armó: {s['usuario']}", f"{s['recogidas']}/{s['unidades']} unidades"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Ubicación", "Artículo", "Color", "Talle", "Cantidad", "Recogido", "Falta", "Código de barras",
               "SKU (ML/TN)", "Publicación", "Ventas / envíos"])
    for cel in ws[3]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="3B1F6B")
    for i in sorted(r.items, key=lambda x: (x.orden_ruta, x.ubicacion, x.articulo, x.color, x.talle)):
        ws.append([i.ubicacion or ("SIN UBICACIÓN" if i.articulo_id else "SIN VINCULAR"), i.articulo, i.color,
                   i.talle, i.cantidad, i.recogido, max(0, i.cantidad - i.recogido), i.codigo_barras, i.sku,
                   i.producto, i.ventas])
        if i.recogido < i.cantidad:
            ws.cell(ws.max_row, 7).font = Font(bold=True, color="C8283E")
    for col, ancho in zip("ABCDEFGHIJK", (14, 22, 18, 8, 10, 10, 8, 16, 30, 40, 30)):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="recoleccion_RC-{r.id}.xlsx"'})

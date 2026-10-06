"""Cierre diario: resumen del día en Excel + mail a gerencia."""
from __future__ import annotations

import io
import smtplib
from collections import Counter, defaultdict
from datetime import date, datetime, time
from email.message import EmailMessage
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select, or_, and_
from sqlalchemy.orm import Session

from .. import config
from ..actividad import ACCIONES
from ..db import ahora
from ..models import Actividad, Cierre, Lote, Paquete, Seccion, Tarea, TIPOS_LOTE


def datos_cierre(db: Session, fecha: date) -> dict:
    inicio, fin = datetime.combine(fecha, time.min), datetime.combine(fecha, time.max)

    # ── Despacho: paquetes de lotes del día + los que quedaron pendientes de días anteriores sin cerrar
    lotes = list(db.scalars(select(Lote).where(or_(Lote.fecha == fecha,
                                                    and_(Lote.fecha < fecha, Lote.cerrado.is_(False))))
                            .order_by(Lote.fecha, Lote.id)).all())
    por_tipo: dict[str, dict] = defaultdict(lambda: {"despachados": 0, "pendientes": 0})
    despachados, pendientes = [], []
    for l in lotes:
        for p in l.paquetes:
            if p.estado == "despachado" and p.escaneado_at and inicio <= p.escaneado_at <= fin:
                despachados.append(p)
                por_tipo[l.tipo]["despachados"] += 1
            elif p.estado == "pendiente":
                pendientes.append(p)
                por_tipo[l.tipo]["pendientes"] += 1

    # ── Tareas por sección
    secciones = {s.slug: s.nombre for s in db.scalars(select(Seccion)).all()}
    tareas = list(db.scalars(select(Tarea).where(or_(Tarea.estado != "hecha",
                                                      and_(Tarea.hecha_at >= inicio, Tarea.hecha_at <= fin)))).all())
    tareas_sec: dict[str, dict] = defaultdict(lambda: {"hechas": 0, "pendientes": 0, "vencidas": 0})
    for t in tareas:
        d = tareas_sec[t.seccion_slug]
        if t.estado == "hecha":
            d["hechas"] += 1
        else:
            d["pendientes"] += 1
            if t.vencida(fecha):
                d["vencidas"] += 1

    # ── Actividad por operario
    acts = list(db.scalars(select(Actividad).where(Actividad.ts >= inicio, Actividad.ts <= fin)
                           .order_by(Actividad.ts)).all())
    por_op: dict[str, Counter] = defaultdict(Counter)
    for a in acts:
        nombre = a.usuario.nombre if a.usuario else "(sistema)"
        por_op[nombre][a.accion] += 1

    return dict(fecha=fecha, lotes=lotes, por_tipo=dict(por_tipo), despachados=despachados,
                pendientes=pendientes, tareas=tareas, tareas_sec=dict(tareas_sec), secciones=secciones,
                actividad=acts, por_operario={k: dict(v) for k, v in por_op.items()})


def _hoja(wb: Workbook, titulo: str, encabezados: list[str], filas: list[list], anchos: list[int]):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


def excel_cierre(d: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Cierre PARKA Depósito — {d['fecha'].strftime('%d/%m/%Y')}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Despacho", "Despachados", "No despachados"])
    for tipo, v in d["por_tipo"].items():
        ws.append([TIPOS_LOTE.get(tipo, tipo), v["despachados"], v["pendientes"]])
    ws.append(["TOTAL", len(d["despachados"]), len(d["pendientes"])])
    ws.append([])
    ws.append(["Sección", "Tareas hechas", "Pendientes", "Vencidas"])
    for slug, v in d["tareas_sec"].items():
        ws.append([d["secciones"].get(slug, slug), v["hechas"], v["pendientes"], v["vencidas"]])
    ws.append([])
    ws.append(["Operario", "Paquetes despachados", "Tareas completadas", "Acciones totales"])
    for op, c in d["por_operario"].items():
        ws.append([op, c.get("escaneo_ok", 0) + c.get("despacho_manual", 0), c.get("tarea_hecha", 0), sum(c.values())])
    for col, w in zip("ABCD", [28, 20, 20, 18]):
        ws.column_dimensions[col].width = w

    def fila_paq(p: Paquete) -> list:
        return [TIPOS_LOTE.get(p.lote.tipo, p.lote.tipo), p.lote.nombre, p.tracking, p.venta_id, p.comprador,
                p.producto, p.sku, p.cantidad,
                p.escaneado_at.strftime("%H:%M") if p.escaneado_at else "",
                p.escaneado_por.nombre if p.escaneado_por else ""]

    enc = ["Tipo", "Lote", "Tracking", "Venta", "Comprador", "Producto", "SKU", "Cant.", "Hora", "Operario"]
    anchos = [12, 26, 18, 18, 24, 40, 18, 7, 8, 14]
    _hoja(wb, "Despachados", enc, [fila_paq(p) for p in d["despachados"]], anchos)
    _hoja(wb, "No despachados", enc[:8], [fila_paq(p)[:8] for p in d["pendientes"]], anchos[:8])
    _hoja(wb, "Tareas", ["Sección", "Tarea", "Origen", "Estado", "Asignada a", "Vence", "Hecha por", "Hora"],
          [[d["secciones"].get(t.seccion_slug, t.seccion_slug), t.titulo, t.origen,
            "Vencida" if t.vencida(d["fecha"]) else t.estado,
            t.asignado_a.nombre if t.asignado_a else "", t.vence.strftime("%d/%m") if t.vence else "",
            t.hecha_por.nombre if t.hecha_por else "", t.hecha_at.strftime("%H:%M") if t.hecha_at else ""]
           for t in d["tareas"]], [16, 50, 8, 10, 14, 8, 14, 8])
    _hoja(wb, "Actividad", ["Hora", "Operario", "Sección", "Acción", "Detalle"],
          [[a.ts.strftime("%H:%M:%S"), a.usuario.nombre if a.usuario else "", a.seccion_slug,
            ACCIONES.get(a.accion, a.accion), a.detalle] for a in d["actividad"]], [10, 14, 12, 24, 60])
    st = _hoja(wb, "Stock y faltantes", ["Concepto", "Detalle"],
               [["Paquetes sin despachar", len(d["pendientes"])],
                ["Diferencias de conteo", "Disponible cuando se active la sección Depósito (etapa 2)"],
                ["Faltantes para despachar", "Disponible en etapa 2"],
                ["Devoluciones reingresadas", "Disponible en etapa 2"]], [30, 60])
    st.sheet_state = "visible"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def html_cierre(d: dict) -> str:
    def tabla(enc: list[str], filas: list[list]) -> str:
        th = "".join(f"<th style='text-align:left;padding:6px 10px;background:#1f2937;color:#fff'>{escape(h)}</th>" for h in enc)
        trs = "".join("<tr>" + "".join(f"<td style='padding:6px 10px;border-bottom:1px solid #e5e7eb'>{escape(str(c))}</td>" for c in f) + "</tr>" for f in filas)
        return f"<table style='border-collapse:collapse;font-size:14px;margin:8px 0 20px'>{'<tr>'+th+'</tr>'}{trs}</table>"

    desp = [[TIPOS_LOTE.get(t, t), v["despachados"], v["pendientes"]] for t, v in d["por_tipo"].items()]
    desp.append(["TOTAL", len(d["despachados"]), len(d["pendientes"])])
    tar = [[d["secciones"].get(s, s), v["hechas"], v["pendientes"], v["vencidas"]] for s, v in d["tareas_sec"].items()]
    ops = [[op, c.get("escaneo_ok", 0) + c.get("despacho_manual", 0), c.get("tarea_hecha", 0), sum(c.values())]
           for op, c in d["por_operario"].items()]
    pend = [[TIPOS_LOTE.get(p.lote.tipo, p.lote.tipo), p.tracking, p.comprador, p.producto[:60]] for p in d["pendientes"][:50]]
    extra = f"<p style='color:#6b7280'>…y {len(d['pendientes']) - 50} más (ver Excel adjunto)</p>" if len(d["pendientes"]) > 50 else ""
    return f"""<div style="font-family:Arial,sans-serif;color:#111">
<h2 style="margin:0 0 4px">Cierre del día — {d['fecha'].strftime('%d/%m/%Y')}</h2>
<p style="color:#6b7280;margin:0 0 16px">PARKA Depósito · generado {ahora().strftime('%H:%M')}</p>
<h3>Despacho</h3>{tabla(['Tipo','Despachados','No despachados'], desp)}
<h3>Tareas por sección</h3>{tabla(['Sección','Hechas','Pendientes','Vencidas'], tar) if tar else '<p>Sin tareas.</p>'}
<h3>Actividad por operario</h3>{tabla(['Operario','Paquetes','Tareas','Acciones'], ops) if ops else '<p>Sin actividad.</p>'}
<h3>No despachados</h3>{tabla(['Tipo','Tracking','Comprador','Producto'], pend) if pend else '<p>Todo despachado ✔</p>'}{extra}
<h3>Stock y faltantes</h3><p>Paquetes sin despachar: <b>{len(d['pendientes'])}</b>. Conteos, faltantes y devoluciones se suman con la sección Depósito (etapa 2).</p>
<p style="color:#6b7280">El detalle completo está en el Excel adjunto.</p></div>"""


def mail_configurado() -> bool:
    return bool(config.SMTP_HOST and config.SMTP_USER and config.SMTP_PASS)


def enviar_cierre(db: Session, fecha: date, automatico: bool = False, destinatario: str | None = None) -> Cierre:
    destinatario = destinatario or config.MAIL_TO
    d = datos_cierre(db, fecha)
    registro = Cierre(fecha=fecha, destinatario=destinatario, automatico=automatico)
    try:
        if not mail_configurado():
            raise RuntimeError("Falta configurar SMTP_USER / SMTP_PASS en el archivo .env")
        msg = EmailMessage()
        msg["Subject"] = f"Cierre PARKA Depósito — {fecha.strftime('%d/%m/%Y')} · {len(d['despachados'])} despachados, {len(d['pendientes'])} pendientes"
        msg["From"] = config.MAIL_FROM or config.SMTP_USER
        msg["To"] = destinatario
        msg.set_content("Tu cliente de correo no muestra HTML. El detalle está en el Excel adjunto.")
        msg.add_alternative(html_cierre(d), subtype="html")
        msg.add_attachment(excel_cierre(d), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=f"cierre_{fecha.isoformat()}.xlsx")
        with smtplib.SMTP(config.SMTP_HOST, config.SMTP_PORT, timeout=30) as s:
            s.starttls()
            s.login(config.SMTP_USER, config.SMTP_PASS)
            s.send_message(msg)
        registro.ok = True
    except Exception as e:  # noqa: BLE001 — se guarda el error para mostrarlo en pantalla
        registro.ok = False
        registro.error = str(e)[:1000]
    db.add(registro)
    db.commit()
    return registro

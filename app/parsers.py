"""Lectura de archivos de despacho — portado del scanner (index.html).

Soporta:
  • Etiquetas ZPL de Mercado Libre (uno o varios bloques ^XA…^XZ)
  • Planillas CSV / XLSX (export de ML o de Tienda Nube) con detección de columnas
"""
from __future__ import annotations

import csv
import io
import json
import re

# ─────────────────────────── ZPL ───────────────────────────

def _decode_zpl(s: str) -> str:
    if not s:
        return ""
    raw = re.sub(r"_([0-9A-Fa-f]{2})", lambda m: chr(int(m.group(1), 16)), s)
    # Los bytes UTF-8 quedaron como caracteres latin-1: reconvertir si se puede
    try:
        raw = raw.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return raw.strip()


def es_zpl(texto: str) -> bool:
    return bool(re.search(r"\^XA", texto, re.I))


_PREFIJOS_NO_TITULO = ("Color:", "SKU:", "Talle:", "Remitente", "Despachar", "Recort", "Pack ID", "Venta",
                       "Domicilio", "Ciudad", "Referencia", "CP:", "Barrio", "Entre")


def _atributos(linea: str) -> dict[str, str] | None:
    """'Color: Negro  | Talle: XL  | SKU: M-114-BLACK-XL' → {'Color': 'Negro', 'Talle': 'XL', 'SKU': '...'}"""
    if ":" not in linea or ("SKU:" not in linea and not re.match(r"^(Color|Talle|Diseño|Nombre del diseño)\s*:", linea)):
        return None
    out = {}
    for parte in linea.split("|"):
        if ":" in parte:
            k, v = parte.split(":", 1)
            if k.strip() and v.strip():
                out[k.strip()] = v.strip()
    return out or None


def items_zpl(block: str) -> tuple[list[dict], int]:
    """Artículos que lleva una etiqueta de ML: título de la publicación + línea de atributos
    (Color / Talle / SKU). Devuelve (items, unidades que dice la etiqueta)."""
    unidades = 0
    m = re.search(r"\^FD(\d{1,4})\^FS\s*\^FO[^\n]*?\^FDUnidad(?:es)?\^FS", block, re.I)
    if m:
        unidades = int(m.group(1))
    items, titulo = [], ""
    for crudo in re.findall(r"\^FD([^\^]*)\^FS", block):
        v = _decode_zpl(crudo)
        attrs = _atributos(v)
        if attrs:
            variante = " · ".join(f"{k}: {x}" for k, x in attrs.items() if k.upper() != "SKU")
            items.append(dict(sku=attrs.get("SKU", ""), producto=titulo, color=attrs.get("Color", ""),
                              talle=attrs.get("Talle", ""), variante=variante, cantidad=1))
            titulo = ""
        elif len(v) >= 10 and not v.startswith(_PREFIJOS_NO_TITULO) and not re.fullmatch(r"[\d\s>:]+", v):
            titulo = v
    if len(items) == 1 and unidades > 1:
        items[0]["cantidad"] = unidades
    return items, unidades or sum(i["cantidad"] for i in items)


def parse_zpl(texto: str) -> list[dict]:
    """Devuelve la lista de paquetes. Lanza ValueError si hay etiquetas sin tracking."""
    texto = texto.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    bloques = re.findall(r"\^XA([\s\S]*?)\^XZ", texto, re.I)
    paquetes: list[dict] = []
    for block in bloques:
        tracking = ""
        # A: Code128 con cualquier orientación/parámetros
        m = re.search(r"\^BC[A-Z](?:,[^,\^]*){0,5}\^FD[^0-9\^]*([0-9]{6,15})\^FS", block, re.I)
        if m:
            tracking = m.group(1)
        # B: QR JSON {"id":"NUMERO","t":"lm"}
        if not tracking:
            m = re.search(r'"id"\s*:\s*"([0-9]{6,15})"', block)
            if m:
                tracking = m.group(1)
        # C: prefijo AIM >: o >8
        if not tracking:
            m = re.search(r"\^FD>[:8]([0-9]{6,15})\^FS", block)
            if m:
                tracking = m.group(1)
        # D: cualquier ^FD con 8-15 dígitos que no sea un ID de venta
        if not tracking:
            for n in re.finditer(r"\^FD([0-9]{8,15})\^FS", block):
                if not n.group(1).startswith("2000"):
                    tracking = n.group(1)
                    break
        if not tracking:
            continue

        venta = ""
        m = re.search(r"\^FD(?:Venta:\s*)?20000\^FS[\s\S]*?\^FD([0-9]{11})\^FS", block, re.I)
        if m:
            venta = "20000" + m.group(1)

        producto = ""
        for p in re.finditer(r"\^FB\d+,\d+[^)]*?\^FH\^FD([^\^]{10,})\^FS", block):
            v = _decode_zpl(p.group(1))
            if not v.startswith(("Color:", "SKU:", "Remitente", "Despachar", "Recort")):
                producto = v
                break

        sku = ""
        m = re.search(r"SKU:\s*([A-Za-z0-9_\-]+)", block)
        if m:
            sku = _decode_zpl(m.group(1))

        comprador = ""
        m = re.search(r"\^FH\^FD([\w\xC0-\xFF][\w\xC0-\xFF\s]{3,50}?)\s*\([A-Z0-9_]{4,}\)\^FS", block)
        if m:
            comprador = _decode_zpl(m.group(1))

        items, unidades = items_zpl(block)
        paquetes.append(dict(tracking=tracking, venta_id=venta, comprador=comprador,
                             producto=producto or sku, sku=sku, cantidad=1, items=items, unidades=unidades))

    if bloques and not paquetes:
        raise ValueError(f"Se encontraron {len(bloques)} etiquetas ZPL pero ninguna tiene número de envío")
    return paquetes


# ─────────────────────── CSV / XLSX ───────────────────────

def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def leer_tabla(nombre: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """Devuelve (encabezados, filas) de un CSV o XLSX."""
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        texto = _decode_bytes(data)
        muestra = texto[:5000]
        try:
            dialect = csv.Sniffer().sniff(muestra, delimiters=",;\t")
            delim = dialect.delimiter
        except csv.Error:
            delim = ";" if muestra.count(";") > muestra.count(",") else ","
        filas = [[c.strip() for c in r] for r in csv.reader(io.StringIO(texto), delimiter=delim)]
    # saltar filas vacías al principio (algunos exports de ML traen títulos arriba)
    while filas and not any(filas[0]):
        filas.pop(0)
    if not filas:
        raise ValueError("El archivo está vacío")
    # si la primera fila tiene 1 sola celda llena, es un título: buscar la fila de encabezados
    for i, f in enumerate(filas[:10]):
        if sum(1 for c in f if c) >= 3:
            filas = filas[i:]
            break
    headers = [h.strip() for h in filas[0]]
    datos = [f for f in filas[1:] if any(c for c in f)]
    if not datos:
        raise ValueError("El archivo no contiene filas de datos")
    return headers, datos


_KW = {
    "tracking": ["envío", "envio", "tracking", "seguimiento", "shipment", "n° envío", "nro envio", "ship_id",
                 "código de envío", "codigo de envio"],
    "venta_id": ["n° venta", "n° de venta", "# de venta", "venta", "sale", "order_id", "id venta",
                 "número de orden", "numero de orden", "orden"],
    "comprador": ["comprador", "buyer", "nombre del comprador", "cliente", "customer", "nombre"],
    "producto": ["título", "titulo", "nombre del producto", "producto", "artículo", "articulo", "product",
                 "item", "description"],
    "sku": ["sku"],
    "cantidad": ["cantidad", "unidades", "quantity", "qty"],
    "variante": ["variante", "variación", "variacion", "variation"],
}


def autodetectar(headers: list[str], tipo: str = "") -> dict[str, str]:
    lc = [h.lower() for h in headers]

    def buscar(kws: list[str], excluir: set[str]) -> str:
        for kw in kws:
            for i, h in enumerate(lc):
                if kw in h and headers[i] not in excluir:
                    return headers[i]
        return ""

    usados: set[str] = set()
    res: dict[str, str] = {}
    for campo in ("tracking", "venta_id", "sku", "cantidad", "variante", "comprador", "producto"):
        res[campo] = buscar(_KW[campo], usados)
        if res[campo]:
            usados.add(res[campo])
    # Tienda Nube: si no hay columna de envío, se escanea el número de orden
    if not res["tracking"] and tipo == "tiendanube" and res["venta_id"]:
        res["tracking"] = res["venta_id"]
    return res


def construir_paquetes(headers: list[str], filas: list[list[str]], mapeo: dict[str, str]) -> list[dict]:
    """Arma paquetes agrupando por tracking (una orden de TN trae una fila por producto)."""
    idx = {c: (headers.index(mapeo[c]) if mapeo.get(c) in headers else -1) for c in _KW}
    if idx["tracking"] < 0:
        raise ValueError("Elegí qué columna tiene el número de envío / código a escanear")

    def val(fila: list[str], campo: str) -> str:
        i = idx[campo]
        return fila[i].strip() if 0 <= i < len(fila) else ""

    por_tracking: dict[str, dict] = {}
    for fila in filas:
        trk = normalizar_codigo(val(fila, "tracking"))
        if not trk:
            continue
        try:
            cant = int(float(val(fila, "cantidad").replace(",", ".") or 1))
        except ValueError:
            cant = 1
        prod, sku = val(fila, "producto"), val(fila, "sku")
        item = dict(sku=sku, producto=prod, color="", talle="", variante=val(fila, "variante"), cantidad=cant)
        if trk in por_tracking:
            p = por_tracking[trk]
            p["items"].append(item)
            p["cantidad"] += cant
            if prod and prod not in p["producto"]:
                p["producto"] = (p["producto"] + " | " + prod).strip(" |")
            if sku and sku not in p["sku"]:
                p["sku"] = (p["sku"] + ", " + sku).strip(", ")
        else:
            por_tracking[trk] = dict(tracking=trk, venta_id=val(fila, "venta_id"),
                                     comprador=val(fila, "comprador"), producto=prod or sku,
                                     sku=sku, cantidad=cant, items=[item])
    return list(por_tracking.values())


def _norm(x: str) -> str:
    """Normaliza encabezados ignorando acentos/mojibake (Número, N�mero -> nmero)."""
    return re.sub(r"[^a-z0-9]", "", str(x).lower())


def es_tiendanube(headers: list[str]) -> bool:
    j = "|".join(h.strip().lower() for h in headers)
    return "mero de orden" in j and "tracking" in j


def construir_tiendanube(headers: list[str], filas: list[list[str]]) -> list[dict]:
    """Export de órdenes de Tienda Nube: una fila por producto → un paquete por orden.
    Se escanea el nº de orden; el código de tracking del envío queda como alternativo."""
    h = [x.strip() for x in headers]

    def idx(nombre: str) -> int:
        n = _norm(nombre)
        return next((i for i, c in enumerate(h) if _norm(c) == n), -1)

    iO, iCE, iC = idx("Número de orden"), idx("Nombre para el envío"), idx("Nombre del comprador")
    iSku, iCant, iProd = idx("SKU"), idx("Cantidad del producto"), idx("Nombre del producto")
    iTrk = idx("Código de tracking del envío")
    iEstEnv, iEstOrd, iMedio = idx("Estado del envío"), idx("Estado de la orden"), idx("Medio de envío")

    def v(row, i):
        return row[i].strip() if 0 <= i < len(row) and row[i] else ""

    orden: dict[str, dict] = {}
    for row in filas:
        o = normalizar_codigo(v(row, iO))
        if not o:
            continue
        g = orden.setdefault(o, dict(tracking=o, venta_id=o, comprador="", productos=[], skus=[],
                                     cantidad=0, tracking_alt="", items=[], estado_envio="", estado_orden="",
                                     medio=""))
        for campo, i in (("estado_envio", iEstEnv), ("estado_orden", iEstOrd), ("medio", iMedio)):
            if not g[campo] and v(row, i):
                g[campo] = v(row, i)
        comp = v(row, iCE) if v(row, iCE) and v(row, iCE).lower() != "no informado" else v(row, iC)
        if not g["comprador"] and comp:
            g["comprador"] = comp
        sku, prod = v(row, iSku), v(row, iProd)
        if sku and sku not in g["skus"]:
            g["skus"].append(sku)
        if prod and prod not in g["productos"]:
            g["productos"].append(prod)
        try:
            cant = int(float(v(row, iCant) or 0))
        except ValueError:
            cant = 0
        g["cantidad"] += cant
        if sku or prod:
            color, talle = variante_tn(prod)
            g["items"].append(dict(sku=sku, producto=prod, color=color, talle=talle,
                                   variante=" · ".join(x for x in (f"Color: {color}" if color else "",
                                                                   f"Talle: {talle}" if talle else "") if x),
                                   cantidad=max(1, cant)))
        trk = re.sub(r'[="]', "", v(row, iTrk)).strip()
        if trk and not g["tracking_alt"]:
            g["tracking_alt"] = trk
    out = []
    for g in orden.values():
        out.append(dict(tracking=g["tracking"], venta_id=g["venta_id"], comprador=g["comprador"],
                        producto=" + ".join(g["productos"]), sku=", ".join(g["skus"]),
                        cantidad=g["cantidad"] or 1, tracking_alt=g["tracking_alt"], items=g["items"],
                        estado_envio=g["estado_envio"], estado_orden=g["estado_orden"], medio=g["medio"],
                        tipo="tiendanube"))
    out.sort(key=lambda p: int(p["tracking"]) if p["tracking"].isdigit() else 0)
    return out


def variante_tn(producto: str) -> tuple[str, str]:
    """"Campera Beast (Gris, M)" → ("Gris", "M"). Sin paréntesis → ("", "")."""
    m = re.search(r"\(([^()]*)\)\s*$", producto or "")
    if not m:
        return "", ""
    partes = [x.strip() for x in m.group(1).split(",") if x.strip()]
    if len(partes) >= 2:
        return ", ".join(partes[:-1]), partes[-1]
    return (partes[0], "") if partes else ("", "")


def tn_pendiente(p: dict) -> bool:
    """Orden de Tienda Nube que todavía hay que armar: no cancelada y sin enviar/entregar."""
    est = (p.get("estado_envio") or "").lower()
    orden = (p.get("estado_orden") or "").lower()
    if "cancel" in orden or "archiv" in orden:
        return False
    return not any(k in est for k in ("enviad", "entregad", "retirad", "despachad"))


# ─────────────────────── Excel de ventas de Mercado Libre ───────────────────────
# "YYYYMMDD_Ventas_AR_Mercado_Libre_y_Mercado_Shops_….xlsx": avisos arriba, una fila de grupos
# (Ventas | Publicidad | Publicaciones | … | Envíos | Devoluciones | Reclamos) y abajo las columnas.
# Hay columnas con el mismo nombre en distintos grupos (Estado, Unidades, Forma de entrega…).

_ML_NO_ARMAR = ("cancel", "entregad", "en camino", "devuel", "no entregad", "mediación")


def _filas_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    filas = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return filas


def tracking_ml(raw: str) -> tuple[str, str]:
    """Nº de seguimiento del Excel → (nº de envío que se escanea, texto original).
    Flex: "48185193093". Colecta: "MEL48183013451FMXDF01" → "48183013451"."""
    s = normalizar_codigo(raw).strip()
    m = re.match(r"^[A-Z]{2,4}(\d{9,})[A-Z0-9]*$", s, re.I)
    if m:
        return m.group(1), s
    return s, ""


def tipo_envio_ml(forma: str) -> str:
    f = (forma or "").lower()
    if "flex" in f:
        return "flex"
    if "colecta" in f or "recolec" in f:
        return "colecta"
    if "full" in f:
        return "full"
    return "ml"


def ventas_ml(nombre: str, data: bytes) -> list[dict] | None:
    """Lee el Excel de ventas de ML. Devuelve paquetes (por nº de envío) o None si no es ese formato."""
    if not nombre.lower().endswith((".xlsx", ".xlsm")):
        return None
    try:
        filas = _filas_xlsx(data)
    except Exception:  # noqa: BLE001
        return None
    fila_cols = next((i for i, f in enumerate(filas[:15])
                      if any(c.strip().lower() == "# de venta" for c in f) and any(c.strip().upper() == "SKU" for c in f)),
                     None)
    if fila_cols is None:
        return None
    cols = [c.strip() for c in filas[fila_cols]]
    grupos, actual = [], ""
    fila_g = filas[fila_cols - 1] if fila_cols else []
    for i in range(len(cols)):
        g = fila_g[i].strip() if i < len(fila_g) else ""
        actual = g or actual
        grupos.append(actual.lower())

    def col(nombre_col: str, grupo: str = "") -> int:
        n = _norm(nombre_col)
        for i, c in enumerate(cols):
            if _norm(c) == n and (not grupo or grupos[i].startswith(grupo)):
                return i
        for i, c in enumerate(cols):  # sin grupo (por si ML cambia los títulos de arriba)
            if _norm(c) == n:
                return i
        return -1

    iV, iEst, iU, iPack = col("# de venta", "venta"), col("Estado", "venta"), col("Unidades", "venta"), \
        col("Paquete de varios productos", "venta")
    iSku, iTit, iVar = col("SKU", "publica"), col("Título de la publicación", "publica"), col("Variante", "publica")
    iComp, iForma = col("Comprador", "compra"), col("Forma de entrega", "env")
    iTrk = col("Número de seguimiento", "env")
    if iTrk < 0:  # sin nº de envío: se lee como planilla común y se cruza con las etiquetas por nº de venta
        return None

    def v(f, i):
        return f[i].strip() if 0 <= i < len(f) and f[i] else ""

    paquetes: dict[str, dict] = {}
    ultimo: dict | None = None
    for f in filas[fila_cols + 1:]:
        if not any(f):
            continue
        venta = normalizar_codigo(v(f, iV))
        if not venta or not venta[:1].isdigit():
            continue
        estado = v(f, iEst)
        trk, alt = tracking_ml(v(f, iTrk))
        sku, titulo = v(f, iSku), v(f, iTit)
        # Paquete de varios productos: la fila del paquete trae el envío; las de cada producto, el SKU
        if not trk and ultimo is not None and (sku or titulo):
            p = ultimo
        else:
            if any(k in estado.lower() for k in _ML_NO_ARMAR):
                ultimo = None
                continue
            if not trk:
                trk = venta  # sin envío (ej. acordar con el comprador): se escanea el nº de venta
            p = paquetes.setdefault(trk, dict(tracking=trk, tracking_alt=alt, venta_id=venta,
                                              comprador=v(f, iComp), items=[], tipo=tipo_envio_ml(v(f, iForma)),
                                              estado=estado))
            ultimo = p
        if not (sku or titulo):
            continue
        try:
            cant = max(1, int(float(v(f, iU).replace(",", ".") or 1)))
        except ValueError:
            cant = 1
        attrs = [a.split(":", 1) for a in v(f, iVar).split("|") if ":" in a]
        attrs = [(k.strip(), x.strip()) for k, x in attrs]
        color = next((x for k, x in attrs if k.lower() == "color"), "")
        talle = next((x for k, x in attrs if k.lower() == "talle"), "")
        p["items"].append(dict(sku=sku, producto=titulo, color=color, talle=talle, cantidad=cant,
                               variante=" · ".join(f"{k}: {x}" for k, x in attrs)))
    return [p for p in paquetes.values() if p["items"]]


def leer_archivo(nombre: str, data: bytes, tipo: str = "", mapeo: dict | None = None) -> dict:
    """Punto de entrada único. Devuelve:
    {"formato": "zpl"|"tabla", "paquetes": [...], "headers": [...], "mapeo": {...}, "muestra": [...]}
    """
    es_texto = not nombre.lower().endswith((".xlsx", ".xlsm"))
    if es_texto:
        texto = _decode_bytes(data)
        if es_zpl(texto):
            return {"formato": "zpl", "paquetes": parse_zpl(texto), "headers": [], "mapeo": {}, "muestra": []}
    headers, filas = leer_tabla(nombre, data)
    if es_tiendanube(headers) and not mapeo:
        return {"formato": "tiendanube", "paquetes": construir_tiendanube(headers, filas), "headers": [],
                "mapeo": {}, "muestra": []}
    mapeo_final = autodetectar(headers, tipo)
    if mapeo:
        mapeo_final.update({k: v for k, v in mapeo.items() if k in _KW and v is not None})
    try:
        paquetes = construir_paquetes(headers, filas, mapeo_final)
    except ValueError:
        paquetes = []
    return {"formato": "tabla", "paquetes": paquetes, "headers": headers,
            "mapeo": mapeo_final, "muestra": filas[:5]}


# ─────────────────────── ESCANEO ───────────────────────

def codigo_comparable(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def normalizar_codigo(s: str) -> str:
    s = (s or "").strip()
    if s.endswith(".0") and s[:-2].isdigit():  # números leídos como float en Excel
        s = s[:-2]
    return s.lstrip("#").strip()


def extraer_id(raw: str) -> str:
    """Saca el número de envío de lo que devuelva el lector:
    QR  → '{"id":"47383333425","t":"lm"}' o 'LA,{…}'
    Barra → '47383333425', ']C047383333425', '>:47383333425'
    """
    s = (raw or "").strip()
    if s.startswith("{"):
        try:
            j = json.loads(s)
            if j.get("id"):
                return str(j["id"])
        except (ValueError, AttributeError):
            pass
    m = re.match(r"^LA[,\s]+(\{.+\})", s, re.I)
    if m:
        try:
            j = json.loads(m.group(1))
            if j.get("id"):
                return str(j["id"])
        except (ValueError, AttributeError):
            pass
    if re.match(r"^\]C\d", s):
        return s[3:].strip()
    if re.match(r"^>[:\d]", s):
        return s[2:].strip()
    return normalizar_codigo(s)


# ─────────────────────── CATÁLOGO DE ARTÍCULOS ───────────────────────

def articulos_desde_pdf(data: bytes) -> list[dict]:
    """Lee el PDF de 'Etiquetas de producto' de Odoo (grilla de 4 columnas).
    Cada etiqueta: Nombre (variante) / SKU / código de barras (algunos productos no tienen código)."""
    from pypdf import PdfReader

    es_ean = re.compile(r"\d{8,14}")
    es_sku = re.compile(r"[A-Za-z0-9][A-Za-z0-9_./]*(?:-[A-Za-z0-9_./]+)+")
    reader = PdfReader(io.BytesIO(data))
    etiquetas: list[dict] = []
    for page in reader.pages:
        items = []

        def visitor(text, cm, tm, fd, fs):
            t = (text or "").strip()
            if t:
                items.append((tm[4], tm[5], t))

        page.extract_text(visitor_text=visitor)
        columnas: dict[int, list] = {}
        for x, y, t in items:
            columnas.setdefault(int((x - 20) // 183), []).append((y, t))
        for col in sorted(columnas):
            actual = None
            ultimo_tipo, ultimo_y = None, -999.0
            for y, t in sorted(columnas[col]):
                if es_ean.fullmatch(t):
                    if actual is None or actual["codigo_barras"]:
                        actual = {"nombre": [], "sku": "", "codigo_barras": ""}
                        etiquetas.append(actual)
                    actual["codigo_barras"] = t
                    tipo = "ean"
                elif " " not in t and es_sku.fullmatch(t) and not t.isdigit():
                    if actual is None or actual["sku"] or actual["codigo_barras"]:
                        actual = {"nombre": [], "sku": "", "codigo_barras": ""}
                        etiquetas.append(actual)
                    actual["sku"] = t
                    tipo = "sku"
                else:
                    continua = ultimo_tipo == "nombre" and (y - ultimo_y) < 18 and actual is not None
                    if not continua:
                        actual = {"nombre": [], "sku": "", "codigo_barras": ""}
                        etiquetas.append(actual)
                    actual["nombre"].append(t)
                    tipo = "nombre"
                ultimo_tipo, ultimo_y = tipo, y
    return _normalizar_articulos([dict(nombre=" ".join(c["nombre"]).strip(), sku=c["sku"],
                                       codigo_barras=c["codigo_barras"]) for c in etiquetas])


def articulos_desde_tabla(nombre: str, data: bytes) -> list[dict]:
    """CSV/XLSX con columnas de código de barras, SKU y nombre (se detectan solas)."""
    headers, filas = leer_tabla(nombre, data)
    lc = [h.lower() for h in headers]

    def col(*kws):
        for kw in kws:
            for i, h in enumerate(lc):
                if kw in h:
                    return i
        return -1

    iE, iS = col("codigo_barras", "código de barras", "codigo de barras", "barcode", "ean"), col("sku", "referencia")
    iN = col("nombre", "producto", "descripcion", "descripción", "name")
    if iE < 0 and iS < 0:
        raise ValueError("No encontré columnas de código de barras ni SKU")

    def v(f, i):
        return normalizar_codigo(f[i]) if 0 <= i < len(f) else ""
    return _normalizar_articulos([dict(codigo_barras=v(f, iE), sku=(f[iS].strip() if 0 <= iS < len(f) else ""),
                                       nombre=(f[iN].strip() if 0 <= iN < len(f) else "")) for f in filas])


def _normalizar_articulos(crudos: list[dict]) -> list[dict]:
    out, vistos = [], set()
    for c in crudos:
        if not c["sku"] and not c["codigo_barras"]:
            continue
        clave = c["codigo_barras"] or c["sku"]
        if clave in vistos:
            continue
        vistos.add(clave)
        nombre = c.get("nombre", "")
        m = re.match(r"^(.*?)\s*\((.*)\)\s*$", nombre)
        articulo, variante = (m.group(1).strip(), m.group(2).strip()) if m else (nombre, "")
        color = talle = ""
        partes = c["sku"].split("-")
        if len(partes) >= 7:  # 1-CAMI-M-00AXE-NVST-XXL-NOBR
            color, talle = partes[4], partes[5]
        out.append(dict(codigo_barras=c["codigo_barras"], sku=c["sku"], nombre=nombre,
                        articulo=articulo, variante=variante, color=color, talle=talle))
    return out

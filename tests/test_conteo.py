import io
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


def _login(c, nombre="Vengil", pin="1234"):
    uid = re.search(r"elegir\((\d+), '" + nombre + r"'\)", c.get("/login").text).group(1)
    assert c.post("/login", data={"usuario_id": uid, "pin": pin}).json()["ok"]


def test_flujo_conteo():
    with TestClient(app) as c:
        _login(c)
        assert "/conteo" in c.get("/").text
        assert c.get("/conteo").status_code == 200
        # incluye No Brand (00AXE) además de Parka y Puffers
        m = c.get("/conteo/api/buscar?q=00axe").json()["modelos"]
        assert m and {v["marca"] for v in m[0]["variantes"]} == {"No Brand"}
        r = c.get("/conteo/api/buscar?q=abba 02003").json()["modelos"]
        abba = next(x for x in r if x["articulo"] == "ABBA 02003")
        assert {"Parka", "No Brand"} <= {v["marca"] for v in abba["variantes"]}
        # escaneo de un código sin marca preselecciona la variante
        r2 = c.get("/conteo/api/buscar?q=2000000035918").json()["modelos"]
        pre = r2[0]["preseleccion"]
        assert pre
        park = next(v for v in abba["variantes"] if v["codigo_barras"] == "2000000035925")
        res = c.post("/conteo/api/confirmar", json={"nota": "Estante A3", "items": [
            {"articulo_id": pre, "cantidad": 5}, {"articulo_id": park["id"], "cantidad": 2},
            {"articulo_id": park["id"], "cantidad": 1}, {"articulo_id": m[0]["variantes"][0]["id"], "cantidad": 0}]}).json()
        ct = res["conteo"]
        assert ct["total"] == 8 and len(ct["items"]) == 2 and ct["nota"] == "Estante A3"
        x = c.get(f"/conteo/{ct['id']}/conteo.xlsx")
        assert x.status_code == 200
        ws = load_workbook(io.BytesIO(x.content)).active
        filas = [r for r in ws.iter_rows(min_row=4, values_only=True)]
        assert filas[-1][-1] == 8 and {f[3] for f in filas[:-1]} == {"Parka", "No Brand"}
        pdf = c.get(f"/conteo/{ct['id']}/etiquetas.pdf?desde=2")
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
        assert c.get("/conteo/api/historial").json()["conteos"][0]["id"] == ct["id"]
        assert c.post("/conteo/api/confirmar", json={"items": []}).status_code == 400
        assert "conteos hoy" in c.get("/").text

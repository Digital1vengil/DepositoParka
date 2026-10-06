import io
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


def _login(c, nombre="Vengil", pin="1234"):
    uid = re.search(r"elegir\((\d+), '" + nombre + r"'\)", c.get("/login").text).group(1)
    assert c.post("/login", data={"usuario_id": uid, "pin": pin}).json()["ok"]


def test_mapeo_y_recoleccion():
    with TestClient(app) as c:
        _login(c)
        inicio = c.get("/").text
        assert "/s/ecommerce" in inicio and "/s/mapeo" in inicio  # Recolección quedó dentro de Recolección Ecommerce
        assert c.get("/mapeo").status_code == 200 and c.get("/recoleccion").status_code == 200

        # ubicaciones: pasillo A, módulos 1-3, 2 niveles; y pasillo B
        assert c.post("/mapeo/api/ubicaciones/lote", json={"pasillo": "a", "modulo_desde": 1, "modulo_hasta": 3,
                                                           "niveles": 2}).json()["creadas"] == 6
        assert c.post("/mapeo/api/ubicaciones/lote", json={"pasillo": "A", "modulo_desde": 1, "modulo_hasta": 3,
                                                           "niveles": 2}).json()["creadas"] == 0
        c.post("/mapeo/api/ubicaciones/lote", json={"pasillo": "B", "modulo_desde": 1, "modulo_hasta": 2, "niveles": 0})
        ubs = {u["codigo"]: u for u in c.get("/mapeo/api/ubicaciones").json()["ubicaciones"]}
        assert {"A-01-1", "A-03-2", "B-01", "B-02"} <= set(ubs)
        assert c.post("/mapeo/api/ubicaciones/reordenar", json={"modo": "zigzag"}).json()["ok"]
        ubs = {u["codigo"]: u for u in c.get("/mapeo/api/ubicaciones").json()["ubicaciones"]}
        assert ubs["B-02"]["orden"] < ubs["B-01"]["orden"]  # vuelta por el pasillo B

        # artículo: un modelo No Brand y otro con varias variantes
        axe = c.get("/mapeo/api/buscar?q=00axe").json()["modelos"][0]
        v1 = axe["variantes"][0]
        r = c.post("/mapeo/api/asignar", json={"ubicacion_id": ubs["A-02-1"]["id"], "modelo": axe["articulo"],
                                                "color": v1["color"]}).json()
        assert r["ok"] and not r.get("repetido")
        assert c.post("/mapeo/api/asignar", json={"ubicacion_id": ubs["A-02-1"]["id"], "modelo": axe["articulo"],
                                                   "color": v1["color"]}).json()["repetido"]
        assert c.get("/mapeo/api/ubicacion?codigo=a-02-1").json()["asignaciones"][0]["modelo"] == axe["articulo"]
        assert c.get("/mapeo/api/buscar?q=00axe").json()["modelos"][0]["ubicaciones"][0]["codigo"] == "A-02-1"
        pdf = c.get(f"/mapeo/etiquetas.pdf?ids={ubs['A-02-1']['id']},{ubs['B-01']['id']}")
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"

        # recolección desde la lista de Despacho: un SKU del catálogo (x2 en dos ventas), uno desconocido
        paquetes = [
            {"tracking": "4400001", "venta_id": "2000011", "producto": "Camisa 00AXE", "sku": v1["sku"], "cantidad": 1},
            {"tracking": "4400002", "venta_id": "2000012", "producto": "Camisa 00AXE", "sku": v1["sku"].lower(), "cantidad": 1},
            {"tracking": "4400003", "venta_id": "2000013", "producto": "Campera THOR 00114", "sku": "ML-THOR-XYZ", "cantidad": 1},
        ]
        rc = c.post("/recoleccion/api/crear/paquetes", json={"paquetes": paquetes, "tipo": "flex"}).json()["recoleccion"]
        assert rc["paquetes"] == 3 and rc["lineas"] == 2 and rc["unidades"] == 3 and rc["sin_vincular"] == 1
        axe_it = next(i for i in rc["items"] if i["articulo_id"])
        assert axe_it["ubicacion"] == "A-02-1" and axe_it["cantidad"] == 2 and "2000012" in axe_it["ventas"]

        # escanear la prenda suma 1; escanear la ubicación avisa qué hay
        r = c.post(f"/recoleccion/api/{rc['id']}/escanear", json={"codigo": v1["codigo_barras"]}).json()
        assert r["item"]["recogido"] == 1 and r["resumen"]["recogidas"] == 1
        r = c.post(f"/recoleccion/api/{rc['id']}/escanear", json={"codigo": "A-02-1"}).json()
        assert r["tipo"] == "ubicacion" and r["items"][0]["id"] == axe_it["id"]
        c.post(f"/recoleccion/api/item/{axe_it['id']}/marcar", json={"todo": True})
        assert c.post(f"/recoleccion/api/{rc['id']}/escanear", json={"codigo": v1["codigo_barras"]}).status_code == 409
        assert c.post(f"/recoleccion/api/{rc['id']}/escanear", json={"codigo": "999999"}).status_code == 404

        # vincular el SKU desconocido: queda recordado para la próxima
        thor = c.get("/recoleccion/api/buscar?q=00114").json()["modelos"][0]["variantes"][0]
        otro = next(i for i in rc["items"] if not i["articulo_id"])
        rc2 = c.post(f"/recoleccion/api/item/{otro['id']}/vincular", json={"articulo_id": thor["id"]}).json()["recoleccion"]
        assert rc2["sin_vincular"] == 0
        assert any(v["sku"] == "ML-THOR-XYZ" for v in c.get("/mapeo/api/vinculos").json()["vinculos"])
        nueva = c.post("/recoleccion/api/crear/paquetes", json={"paquetes": [paquetes[2]]}).json()["recoleccion"]
        assert nueva["sin_vincular"] == 0 and nueva["sin_ubicacion"] == 1

        # después de asignar ubicación al modelo, "Actualizar" la encuentra
        c.post("/mapeo/api/asignar", json={"ubicacion_id": ubs["B-01"]["id"], "modelo": thor["nombre"].split(" (")[0]})
        act = c.post(f"/recoleccion/api/{nueva['id']}/actualizar").json()["recoleccion"]
        assert act["items"][0]["ubicacion"] == "B-01"

        # terminar, excel e historial
        cerr = c.post(f"/recoleccion/api/{rc['id']}/cerrar").json()["recoleccion"]
        assert cerr["estado"] == "cerrada" and cerr["faltan"] == 1
        assert c.post(f"/recoleccion/api/item/{otro['id']}/marcar", json={"delta": 1}).status_code == 400
        x = c.get(f"/recoleccion/{rc['id']}/recoleccion.xlsx")
        ws = load_workbook(io.BytesIO(x.content)).active
        filas = list(ws.iter_rows(min_row=4, values_only=True))
        assert filas[0][0] == "A-02-1" and filas[0][6] == 0 and any(f[6] == 1 for f in filas)
        assert c.get("/recoleccion/api/lista").json()["recolecciones"][0]["id"] == nueva["id"]
        assert "búsqueda" in c.get("/").text  # la tarjeta ahora es Recolección Ecommerce
        assert c.delete(f"/recoleccion/api/{nueva['id']}").json()["ok"]

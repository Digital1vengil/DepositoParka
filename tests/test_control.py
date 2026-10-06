"""Control de paquetes: etiquetas reales de ML (anonimizadas), SKU de ML → catálogo de Odoo,
escaneo de etiqueta + prendas, bloqueo por prenda equivocada, faltantes y planilla de ventas."""
import io
import re
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app

ZPL = (Path(__file__).parent / "datos" / "etiquetas_ml_ejemplo.txt").read_bytes()


def _login(c, nombre="Vengil", pin="1234"):
    uid = re.search(r"elegir\((\d+), '" + nombre + r"'\)", c.get("/login").text).group(1)
    assert c.post("/login", data={"usuario_id": uid, "pin": pin}).json()["ok"]


def test_parser_items_zpl():
    from app.parsers import leer_archivo
    ps = leer_archivo("etiquetas.txt", ZPL)["paquetes"]
    assert len(ps) == 19
    p = ps[0]
    assert p["tracking"] == "48174207710" and p["venta_id"] == "2000015356940353"
    assert p["items"] == [{"sku": "M-114-BLACK-XL", "producto": "Campera Abrigo Puffer Parka Asgard Thor Hombre Impermeable",
                           "color": "Negro", "talle": "XL", "variante": "Color: Negro · Talle: XL", "cantidad": 1}]


def test_sku_ml_a_catalogo():
    from app.db import SessionLocal
    from app.services import mapeo
    db = SessionLocal()
    try:
        a = mapeo.articulo_por_sku(db, "M-114-BLACK-XL")
        assert a.sku == "1-CAMP-M-00114-BLCK-XL-PARK"
        assert {x.sku for x in mapeo.equivalentes(db, a)} == {"1-CAMP-M-00114-BLCK-XL-PARK", "1-CAMP-M-00114-BLCK-XL-NOBR"}
        assert mapeo.articulo_por_sku(db, "W-2055-BEIG-S").sku == "1-CAMP-W-02055-BEIG-S-PARK"
        assert mapeo.articulo_por_sku(db, "W-8-BLK-M").sku.startswith("1-CAMP-W-00008-BLCK-M-")
        assert mapeo.articulo_por_sku(db, "MOCHI-B3001-BLK") is None
        assert mapeo.colores_equivalentes("BLK", "BLCK") and mapeo.colores_equivalentes("GRAY", "GREY")
        assert not mapeo.colores_equivalentes("BLK", "BLUE")
    finally:
        db.close()


def _barcode(c, sku):
    from app.db import SessionLocal
    from app.models import Articulo
    from sqlalchemy import select
    db = SessionLocal()
    try:
        return db.scalar(select(Articulo.codigo_barras).where(Articulo.sku == sku))
    finally:
        db.close()


def test_flujo_control():
    with TestClient(app) as c:
        _login(c)
        assert "/s/ecommerce" in c.get("/").text  # Control quedó dentro de Recolección Ecommerce
        assert c.get("/control").status_code == 200

        r = c.post("/control/api/crear/archivo", data={"tipo": "flex", "nombre": "Flex prueba"},
                   files=[("archivos", ("Etiqueta de envio.txt", ZPL, "text/plain"))]).json()
        cp = r["control"]
        assert cp["paquetes"] == 19 and cp["pendiente"] == 19
        assert cp["sin_vincular"] == 3  # 2 mochilas y 1 bolso no están en el catálogo
        cid = cp["id"]

        # una prenda sin paquete abierto → guía: en qué paquete va
        thor_park = _barcode(c, "1-CAMP-M-00114-BLCK-XL-PARK")
        g = c.post(f"/control/api/{cid}/escanear", json={"codigo": thor_park}).json()
        assert g["tipo"] == "guia" and g["paquetes"][0]["tracking"] == "48174207710"

        # etiqueta del envío (como la lee el lector: QR) → abre el paquete
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": '{"id":"48174207710","t":"lm"}'}).json()
        assert r["tipo"] == "paquete"
        pk = r["paquete"]
        it = pk["items"][0]
        assert it["articulo"] == "THOR 00114" and it["talle"] == "XL" and it["vinculado"]

        # prenda equivocada → no avanza, queda el error
        otra = _barcode(c, "1-CAMP-W-02101-BLCK-M-PARK")
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": otra, "paquete_id": pk["id"]}).json()
        assert r["ok"] is False and r["motivo"] == "equivocado"
        assert r["paquete"]["escaneadas"] == 0 and r["paquete"]["errores"] == 1
        assert "THOR 00114" in r["esperado"][0]

        # código que no existe
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": "999", "paquete_id": pk["id"]}).json()
        assert r["ok"] is False and r["motivo"] == "desconocido"

        # la misma prenda en No Brand también sirve → paquete controlado
        thor_nobr = _barcode(c, "1-CAMP-M-00114-BLCK-XL-NOBR")
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": thor_nobr, "paquete_id": pk["id"]}).json()
        assert r["ok"] and r["completo"] and r["paquete"]["estado"] == "completo"
        assert r["resumen"]["completo"] == 1

        # una más del mismo: el paquete ya está cerrado
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": thor_park, "paquete_id": pk["id"]}).json()
        assert r["ok"] is False and r["motivo"] == "cerrado"

        # Despacho pregunta por el estado del envío
        assert c.get("/control/api/estado?codigo=48174207710").json()["controlado"] is True
        assert c.get("/control/api/estado?codigo=48171251333").json()["estado"] == "pendiente"
        assert c.get("/control/api/estado?codigo=111").json()["en_control"] is False

        # faltante: segundo paquete (ARCADIA) → tarea automática
        r = c.post(f"/control/api/{cid}/escanear", json={"codigo": "48171251333"}).json()
        p2 = r["paquete"]
        assert c.get(f"/control/api/item/{p2['items'][0]['id']}/donde").json()["opciones"]
        r = c.post(f"/control/api/item/{p2['items'][0]['id']}/faltante", json={"nota": "no hay stock"}).json()
        assert r["paquete"]["estado"] == "faltante" and r["resumen"]["faltante"] == 1
        secs = c.get("/tareas/api/lista").json()["secciones"]
        assert any("Faltante para envío 48171251333" in t["titulo"] for s in secs for t in s["tareas"])
        r = c.post(f"/control/api/item/{p2['items'][0]['id']}/quitar-faltante").json()
        assert r["paquete"]["estado"] == "pendiente"  # todavía no se escaneó nada

        # vincular la mochila (sin SKU en el catálogo) a un artículo: queda recordado
        mochila = next(p for p in c.get(f"/control/api/{cid}").json()["control"]["lista"] if "Mochila" in p["resumen"])
        det = c.get(f"/control/api/paquete/{mochila['id']}").json()["paquete"]
        cualquiera = c.get("/control/api/catalogo?q=00axe").json()["modelos"][0]["variantes"][0]["id"]
        r = c.post(f"/control/api/item/{det['items'][0]['id']}/vincular", json={"articulo_id": cualquiera}).json()
        assert r["lineas"] == 2 and r["paquete"]["items"][0]["vinculado"]  # las dos mochilas del día

        # buscador manual
        assert c.get("/control/api/buscar?q=Prueba3").json()["paquetes"][0]["tracking"] == "48171519368"
        assert c.get("/control/api/buscar?q=W-2101-BLK-M").json()["paquetes"]

        # eventos y Excel
        evs = c.get(f"/control/api/{cid}/eventos").json()["eventos"]
        assert {e["tipo"] for e in evs} >= {"equivocado", "desconocido", "faltante"}
        x = c.get(f"/control/{cid}/control.xlsx")
        wb = load_workbook(io.BytesIO(x.content))
        assert wb["Paquetes"]["A4"].value == "48174207710" and wb["Paquetes"]["D4"].value == "Controlado"
        assert wb["Errores y faltantes"].max_row >= 4

        # otro control: escanear un envío de este → avisa en qué control está
        r2 = c.post("/control/api/crear/archivo", data={"tipo": "flex"},
                    files=[("archivos", ("e.txt", ZPL.split(b"^XZ")[0] + b"^XZ", "text/plain"))]).json()["control"]
        resp = c.post(f"/control/api/{r2['id']}/escanear", json={"codigo": "48171251333"})
        assert resp.status_code == 404 and "otro control" in resp.json()["detail"]


def test_planilla_de_ventas_completa_las_etiquetas():
    """Excel de ventas de ML sin nº de envío: se cruza con las etiquetas por nº de venta."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Ventas", "", "", "Publicaciones", "", "", "Compradores"])
    ws.append(["# de venta", "Fecha de venta", "Unidades", "SKU", "Título de la publicación", "Variante", "Comprador",
               "Costo de envío"])
    ws.append(["2000015356940353", "5 oct", 1, "M-114-BLACK-XL", "Campera Thor", "Color : Negro | Talle : XL", "C1", 0])
    ws.append(["2000099999999999", "5 oct", 2, "W-8-BLK-M", "Campera Harmony", "Color : Negro | Talle : M", "C9", 0])
    buf = io.BytesIO()
    wb.save(buf)
    from app.services.control import leer_archivos
    ps, errores = leer_archivos([("etiquetas.txt", ZPL), ("ventas.xlsx", buf.getvalue())])
    assert len(ps) == 19
    p = next(x for x in ps if x["tracking"] == "48174207710")
    assert p["items"][0]["sku"] == "M-114-BLACK-XL" and p["items"][0]["cantidad"] == 1
    assert any("1 ventas de la planilla" in e for e in errores)
    with TestClient(app) as c:
        _login(c)
        cid = c.post("/control/api/crear/archivo", data={"tipo": "flex"},
                     files=[("archivos", ("e.txt", ZPL, "text/plain"))]).json()["control"]["id"]
        r = c.post(f"/control/api/{cid}/agregar",
                   files=[("archivos", ("ventas.xlsx", buf.getvalue(), "application/octet-stream"))]).json()
        assert r["ok"] and r["nuevos"] == 0


def test_varios_tipos_juntos():
    """Flex + Colecta en la misma carga (Drive o archivos)."""
    with TestClient(app) as c:
        _login(c)
        cp = c.post("/control/api/crear/archivo", data={"tipo": "flex,colecta"},
                    files=[("archivos", ("e.txt", ZPL, "text/plain"))]).json()["control"]
        assert cp["nombre"].startswith("Flex + Colecta") and cp["tipo"] == "mixto" and cp["paquetes"] == 19
        rc = c.post("/recoleccion/api/crear/archivo", data={"tipo": "flex,colecta"},
                    files=[("archivos", ("e.txt", ZPL, "text/plain"))]).json()["recoleccion"]
        assert rc["nombre"].startswith("Flex + Colecta") and rc["tipo"] == "mixto"
